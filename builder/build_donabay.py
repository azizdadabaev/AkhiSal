# -*- coding: utf-8 -*-
"""
Builds Donabay.xlsx - daily filler-block production log and weekly crew payroll.

    python3 builder/build_donabay.py            # the clean workbook to use
    python3 builder/build_donabay.py --demo OUT # a copy filled with test data

How the money works
-------------------
Each worker has one running account. Earnings go in; every som handed over
(advance, weekly pay, correction) comes out. The balance is what you owe the
worker - negative means the worker owes you. Nothing is ever "forgotten": a
debt that a week's earnings cannot cover simply stays on the account and shows
up next week. That is the whole carry-forward mechanism.

Weekly earnings for a worker
    pot   = sum over the week of  paid_blocks x rate_on_that_day
    share = pot x worker_days / crew_days
where paid_blocks = moulded - broken (less any breakage allowance) and days
are 1 for a full day, 0.5 for a half day.

Only Excel-2007-era functions are used (SUMIFS, SUMPRODUCT, INDEX, MATCH,
COUNTIFS, IFERROR), so the file works in any Excel version and in LibreOffice.
"""
import sys, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# ------------------------------------------------------------------ sizes
FIRST_DAY  = dt.date(2026, 9, 21)     # a Monday: the week this workbook starts
N_WEEKS    = 53                       # one year of weeks
N_DAYS     = N_WEEKS * 7
N_SLOTS    = 12                       # worker slots; never reused, see Workers
N_LEDGER   = 1000                     # cash ledger rows
N_RATES    = 20
N_SLIP_ADV = 10                       # advance lines listed on a payslip

# ------------------------------------------------------------------ styles
F = "Arial"
H1   = Font(name=F, size=16, bold=True, color="1F3864")
H2   = Font(name=F, size=12, bold=True, color="1F3864")
HDR  = Font(name=F, size=10, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
BIG  = Font(name=F, size=13, bold=True, color="1F3864")
NOTE = Font(name=F, size=9, italic=True, color="595959")
INP  = Font(name=F, size=10, color="0000FF")
CALC = Font(name=F, size=10, color="000000")
LINK = Font(name=F, size=10, color="008000")

HDR_FILL = PatternFill("solid", fgColor="1F3864")
INP_FILL = PatternFill("solid", fgColor="FFF2CC")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
KEY_FILL = PatternFill("solid", fgColor="DDEBF7")
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
WEEKEND_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")

MONEY = '# ##0;-# ##0;"-"'
MONEY_SIGNED = '# ##0;[Red]-# ##0;"-"'
DATE = "yyyy-mm-dd"
PCT = "0.0%"

def q(name):
    return f"'{name}'" if " " in name else name

S_START, S_LOG, S_LEDGER, S_PAY, S_SLIP = "Start here", "Daily Log", "Cash Ledger", "Weekly Pay", "Payslip"
S_HIST, S_WORKERS, S_SET, S_CHECKS, S_WEEKS = "Worker History", "Workers", "Settings", "Checks", "Weeks"
LOG, LED, WKS, WRK, SET = q(S_LOG), q(S_LEDGER), q(S_WEEKS), q(S_WORKERS), q(S_SET)

def put(ws, ref, value, font=BODY, fill=None, fmt=None, align=None, border=False):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if align: c.alignment = align
    if border: c.border = BOX
    return c

def header_row(ws, row, col0, labels, widths=None, height=30):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=label)
        c.font, c.fill, c.border = HDR, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[L(col0 + i)].width = widths[i]
    ws.row_dimensions[row].height = height

def title(ws, text, sub):
    put(ws, "A1", text, H1)
    put(ws, "A2", sub, NOTE)
    ws.row_dimensions[1].height = 24
    ws.sheet_view.showGridLines = False

def name(wb, key, ref):
    wb.defined_names[key] = DefinedName(key, attr_text=ref)

def dv_list(ws, source, rng, prompt=None):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    if prompt:
        dv.prompt, dv.showInputMessage = prompt, True
    dv.error, dv.showErrorMessage = "Pick a value from the list.", True
    ws.add_data_validation(dv)
    dv.add(rng)

def dv_number(ws, rng, lo, hi, msg, decimal=True):
    dv = DataValidation(type="decimal" if decimal else "whole", operator="between",
                        formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.error, dv.showErrorMessage = msg, True
    ws.add_data_validation(dv)
    dv.add(rng)

def dv_date(ws, rng):
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    dv.error, dv.showErrorMessage = "Type a date, e.g. 2026-09-23.", True
    ws.add_data_validation(dv)
    dv.add(rng)

def flag_rule(ws, rng, first_cell):
    """Pink fill whenever a check cell has text in it."""
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEN({first_cell})>0'], fill=WARN_FILL))


def build(demo=False, week=None, cap=None):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    # sheet order is the order the owner meets them in
    ws_start = wb.active; ws_start.title = S_START
    ws_log   = wb.create_sheet(S_LOG)
    ws_led   = wb.create_sheet(S_LEDGER)
    ws_pay   = wb.create_sheet(S_PAY)
    ws_slip  = wb.create_sheet(S_SLIP)
    ws_hist  = wb.create_sheet(S_HIST)
    ws_wrk   = wb.create_sheet(S_WORKERS)
    ws_set   = wb.create_sheet(S_SET)
    ws_chk   = wb.create_sheet(S_CHECKS)
    ws_wks   = wb.create_sheet(S_WEEKS)

    # ============================================================ Settings
    ws = ws_set
    title(ws, "Settings", "Change these rarely. Yellow cells are yours to edit.")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 60

    put(ws, "B4", "Rules", H2)
    rules = [
        (5, "Breakage allowance", 0, PCT,
         "Broken blocks up to this share of what was moulded are still paid. 0% = pay only good "
         "blocks (your rule). Block plants commonly break 2-5%; set e.g. 3% if you want to absorb "
         "normal breakage yourself.", "BreakAllow"),
        (6, "Old debt taken back per week, max", 1, PCT,
         "When a worker owes you from earlier weeks, at most this share of a week's earnings goes "
         "to paying it back. 100% = take back everything you can (your rule). For formally employed "
         "workers Uzbek labour law caps wage deductions at 50% - check the Labour Code on lex.uz "
         "before relying on this.", "DebtCap"),
        (7, "Reject rate that raises a warning", 0.08, PCT,
         "Normal breakage in block plants is 2-5%; above about 8% usually means a curing or "
         "handling problem worth looking at.", "RejectAlarm"),
        (8, "Expected cement bags per 1000 blocks (optional)", None, "0.0",
         "If you log cement use, the week summary compares it with this. A big gap either way "
         "is the classic sign that block counts are off. Leave blank to skip.", "CementNorm"),
    ]
    for row, label, val, fmt, note, key in rules:
        put(ws, f"B{row}", label, BODY)
        put(ws, f"C{row}", val, INP, INP_FILL, fmt, CENTER, True)
        put(ws, f"D{row}", note, NOTE, align=WRAP)
        ws.row_dimensions[row].height = 48
        name(wb, key, f"{SET}!$C${row}")

    put(ws, "B11", "Pay rate per paid block", H2)
    header_row(ws, 12, 2, ["Starts from (date)", "Rate (so'm per block)"], height=30)
    R0, R1 = 13, 13 + N_RATES - 1
    for r in range(R0, R1 + 1):
        put(ws, f"B{r}", None, INP, INP_FILL, DATE, CENTER, True)
        put(ws, f"C{r}", None, INP, INP_FILL, MONEY, CENTER, True)
    ws[f"B{R0}"].value = FIRST_DAY
    ws[f"B{R0}"].comment = Comment(
        "Type YOUR rate next to this date. To change the rate later, add a new row below with the "
        "date the new rate starts - never overwrite an old row, or past weeks get recalculated at "
        "the new rate. Keep the dates in order, oldest first.", "Donabay")
    put(ws, f"D{R0}", "<- type your rate here. Add a row below when it changes; old weeks keep the old rate.",
        NOTE, align=WRAP)
    name(wb, "RateDates", f"{SET}!$B${R0}:$B${R1}")
    name(wb, "RateValues", f"{SET}!$C${R0}:$C${R1}")
    dv_date(ws, f"B{R0}:B{R1}")
    dv_number(ws, f"C{R0}:C{R1}", 0, 1000000, "Type the rate in so'm per block.")

    # fixed lists behind the dropdowns - named, because Excel 2007 will not
    # take a dropdown source from another sheet any other way
    put(ws, "F4", "Lists (used by dropdowns)", NOTE)
    for i, v in enumerate(["Advance", "Weekly pay", "Correction"]):
        put(ws, f"F{5+i}", v, NOTE)
    for i, v in enumerate(["Active", "Left"]):
        put(ws, f"G{5+i}", v, NOTE)
    for i, v in enumerate(["Yes", "No"]):
        put(ws, f"H{5+i}", v, NOTE)
    name(wb, "EntryTypes", f"{SET}!$F$5:$F$7")
    name(wb, "StatusList", f"{SET}!$G$5:$G$6")
    name(wb, "YesNo", f"{SET}!$H$5:$H$6")

    # ============================================================ Workers
    ws = ws_wrk
    title(ws, "Workers", "Add a worker: fill the next empty slot and set Active. Remove one: set Left. "
                         "Never delete or reuse a slot - that keeps each person's history intact.")
    W0, W1 = 5, 5 + N_SLOTS - 1
    header_row(ws, 4, 1, ["ID", "Name", "Status", "Joined", "Left on", "Phone", "Notes", "Check"],
               [7, 24, 10, 12, 12, 16, 30, 34])
    for i in range(N_SLOTS):
        r = W0 + i
        put(ws, f"A{r}", f"W{i+1:02d}", BOLD, CALC_FILL, align=CENTER, border=True)
        put(ws, f"B{r}", None, INP, INP_FILL, border=True)
        put(ws, f"C{r}", None, INP, INP_FILL, align=CENTER, border=True)
        put(ws, f"D{r}", None, INP, INP_FILL, DATE, CENTER, True)
        put(ws, f"E{r}", None, INP, INP_FILL, DATE, CENTER, True)
        put(ws, f"F{r}", None, INP, INP_FILL, "@", border=True)
        put(ws, f"G{r}", None, INP, INP_FILL, border=True)
        put(ws, f"H{r}",
            f'=IF($B{r}="",IF($C{r}<>"","status set but no name",""),'
            f'IF(COUNTIF($B${W0}:$B${W1},$B{r})>1,"same name twice - make it unique",'
            f'IF($C{r}="","choose Active or Left","")))',
            CALC, CALC_FILL, border=True)
    for i in range(4):
        ws[f"B{W0+i}"].value = f"Worker {i+1}"
        ws[f"C{W0+i}"].value = "Active"
        ws[f"D{W0+i}"].value = FIRST_DAY
    ws[f"B{W0}"].comment = Comment(
        "Placeholder names - type the real ones over them. Everything else in the workbook "
        "follows the ID, so renaming here updates every sheet.", "Donabay")
    flag_rule(ws, f"H{W0}:H{W1}", f"H{W0}")
    dv_list(ws, "=StatusList", f"C{W0}:C{W1}")
    dv_date(ws, f"D{W0}:E{W1}")
    name(wb, "WorkerIDs", f"{WRK}!$A${W0}:$A${W1}")
    name(wb, "WorkerNames", f"{WRK}!$B${W0}:$B${W1}")
    ws.freeze_panes = "A5"

    # ============================================================ Daily Log
    ws = ws_log
    title(ws, "Daily Log", "One row per day. Fill: blocks moulded, blocks broken, and who worked "
                           "(1 = full day, 0.5 = half day, empty = absent).")
    D0, D1 = 6, 6 + N_DAYS - 1
    ATT0 = 8                                  # first attendance column (H)
    ATT1 = ATT0 + N_SLOTS - 1                 # last attendance column (S)
    cCREW, cRATE, cVAL, cWEEK, cCHK, cNOTE = (L(ATT1 + k) for k in range(1, 7))
    a0, a1 = L(ATT0), L(ATT1)

    base = ["Date", "Day", "Moulded", "Broken", "Good", "Paid blocks", "Cement bags"]
    tail = ["Crew-days", "Rate", "Pay value", "Week of", "Check", "Notes"]
    header_row(ws, 5, 1, base + [""] * N_SLOTS + tail,
               [11, 5, 10, 9, 9, 10, 9] + [9] * N_SLOTS + [9, 8, 12, 11, 30, 26], height=42)
    for k in range(N_SLOTS):
        col = ATT0 + k
        # row 4: the slot's name, blank when the slot is empty (used by checks)
        put(ws, f"{L(col)}4", f'=IF({WRK}!$B${W0+k}="","",{WRK}!$B${W0+k})', NOTE)
        ws.cell(row=5, column=col).value = f'=IF({WRK}!$B${W0+k}="",{WRK}!$A${W0+k},{WRK}!$B${W0+k})'
    put(ws, "H3", "Attendance: 1 full day, 0.5 half day, empty = absent", NOTE)

    # totals above the data, so appended rows never fall outside them
    put(ws, "A3", "All time", BOLD)
    for col in "CDEF":
        put(ws, f"{col}3", f"=SUM({col}{D0}:{col}{D1})", BOLD, fmt=MONEY)
    put(ws, f"{cVAL}3", f"=SUM({cVAL}{D0}:{cVAL}{D1})", BOLD, fmt=MONEY)

    for i in range(N_DAYS):
        r = D0 + i
        day = FIRST_DAY + dt.timedelta(days=i)
        put(ws, f"A{r}", day, BODY, fmt=DATE, align=CENTER, border=True)
        # CHOOSE, not TEXT(...,"ddd"): TEXT format codes change with the
        # Excel language, and would show the wrong thing in Russian Excel
        put(ws, f"B{r}", f'=CHOOSE(WEEKDAY($A{r},2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun")',
            CALC, align=CENTER, border=True)
        for col in "CDG":
            put(ws, f"{col}{r}", None, INP, INP_FILL, "# ##0", CENTER, True)
        put(ws, f"E{r}", f'=IF($C{r}="","",$C{r}-N($D{r}))', CALC, CALC_FILL, "# ##0", CENTER, True)
        put(ws, f"F{r}", f'=IF($C{r}="","",$C{r}-MAX(0,N($D{r})-BreakAllow*$C{r}))',
            CALC, CALC_FILL, "# ##0", CENTER, True)
        for col in range(ATT0, ATT1 + 1):
            put(ws, f"{L(col)}{r}", None, INP, INP_FILL, "0.#", CENTER, True)
        put(ws, f"{cCREW}{r}", f'=IF(COUNT({a0}{r}:{a1}{r})=0,"",SUM({a0}{r}:{a1}{r}))',
            CALC, CALC_FILL, "0.#", CENTER, True)
        # latest rate whose start date is on or before this day
        put(ws, f"{cRATE}{r}",
            f'=IF($C{r}="","",IFERROR(INDEX(RateValues,MATCH(SUMPRODUCT(MAX((RateDates<=$A{r})*RateDates)),RateDates,0)),0))',
            CALC, CALC_FILL, MONEY, CENTER, True)
        put(ws, f"{cVAL}{r}", f'=IF($C{r}="","",$F{r}*$' + cRATE + f'{r})', CALC, CALC_FILL, MONEY, None, True)
        put(ws, f"{cWEEK}{r}", f"=$A{r}-WEEKDAY($A{r},2)+1", CALC, CALC_FILL, DATE, CENTER, True)
        put(ws, f"{cCHK}{r}",
            f'=IF(AND($C{r}="",{cCREW}{r}=""),"",'
            f'IF(AND(N($C{r})>0,N({cCREW}{r})=0),"nobody marked present",'
            f'IF(N($D{r})>N($C{r}),"more broken than moulded",'
            f'IF(AND($C{r}="",N({cCREW}{r})>0),"attendance marked but no blocks entered",'
            f'IF(AND(N($C{r})>0,N({cRATE}{r})=0),"no pay rate for this date - see Settings",'
            f'IF(SUMPRODUCT(({a0}{r}:{a1}{r}<>"")*({a0}$4:{a1}$4=""))>0,"attendance in an empty worker slot",""))))))',
            CALC, CALC_FILL, border=True)
        put(ws, f"{cNOTE}{r}", None, INP, INP_FILL, border=True)
        if day.weekday() == 6:                # Sunday: shade the date so weeks read easily
            ws[f"A{r}"].fill = WEEKEND_FILL
            ws[f"B{r}"].fill = WEEKEND_FILL
    flag_rule(ws, f"{cCHK}{D0}:{cCHK}{D1}", f"{cCHK}{D0}")
    dv_number(ws, f"C{D0}:D{D1}", 0, 100000, "Type a whole number of blocks.", decimal=False)
    dv_number(ws, f"G{D0}:G{D1}", 0, 10000, "Type the number of cement bags used.")
    dv_number(ws, f"{a0}{D0}:{a1}{D1}", 0, 1, "Use 1 for a full day, 0.5 for a half day, or leave empty.")
    ws.freeze_panes = f"C{D0}"
    today_row = D0 + (dt.date(2026, 9, 23) - FIRST_DAY).days

    # ============================================================ Cash Ledger
    ws = ws_led
    title(ws, "Cash Ledger", "Every som you hand a worker goes here - one row each time. "
                             "Never delete a row: fix a mistake with a Correction.")
    G0, G1 = 6, 6 + N_LEDGER - 1
    header_row(ws, 5, 1, ["#", "Date", "Worker", "Type", "Amount (so'm)", "Reason", "Given by",
                          "Signed?", "ID", "Week of", "key", "Worker balance after", "Check"],
               [5, 11, 20, 12, 14, 26, 12, 8, 6, 11, 4, 16, 44])
    put(ws, "C3", "Totals:", BOLD)
    for col, typ in (("D", "Advance"), ("F", "Weekly pay"), ("H", "Correction")):
        pass
    put(ws, "D3", "Advances", NOTE, align=Alignment(horizontal="right"))
    put(ws, "E3", f'=SUMIFS($E${G0}:$E${G1},$D${G0}:$D${G1},"Advance")', BOLD, fmt=MONEY)
    put(ws, "F3", "Weekly pay", NOTE, align=Alignment(horizontal="right"))
    put(ws, "G3", f'=SUMIFS($E${G0}:$E${G1},$D${G0}:$D${G1},"Weekly pay")', BOLD, fmt=MONEY)
    put(ws, "I3", f'=SUMIFS($E${G0}:$E${G1},$D${G0}:$D${G1},"Correction")', BOLD, fmt=MONEY)
    put(ws, "H3", "Corrections", NOTE, align=Alignment(horizontal="right"))

    for i in range(N_LEDGER):
        r = G0 + i
        put(ws, f"A{r}", f'=IF($B{r}="","",ROW()-{G0-1})', CALC, CALC_FILL, align=CENTER, border=True)
        put(ws, f"B{r}", None, INP, INP_FILL, DATE, CENTER, True)
        put(ws, f"C{r}", None, INP, INP_FILL, border=True)
        put(ws, f"D{r}", None, INP, INP_FILL, align=CENTER, border=True)
        put(ws, f"E{r}", None, INP, INP_FILL, MONEY, border=True)
        put(ws, f"F{r}", None, INP, INP_FILL, border=True)
        put(ws, f"G{r}", None, INP, INP_FILL, border=True)
        put(ws, f"H{r}", None, INP, INP_FILL, align=CENTER, border=True)
        # the maths is keyed on the worker's ID, not their name
        put(ws, f"I{r}", f'=IF($C{r}="","",IFERROR(INDEX(WorkerIDs,MATCH($C{r},WorkerNames,0)),"?"))',
            CALC, CALC_FILL, align=CENTER, border=True)
        put(ws, f"J{r}", f'=IF($B{r}="","",$B{r}-WEEKDAY($B{r},2)+1)', CALC, CALC_FILL, DATE, CENTER, True)
        # "W01|46286|2" = worker W01's 2nd advance in the week starting on day 46286;
        # the payslip looks advances up by this key. Plain numbers, so it does not
        # depend on how the computer formats dates.
        put(ws, f"K{r}",
            f'=IF(OR($I{r}="",$I{r}="?",$D{r}<>"Advance",$B{r}=""),"",'
            f'$I{r}&"|"&$J{r}&"|"&COUNTIFS($I${G0}:$I{r},$I{r},$J${G0}:$J{r},$J{r},$D${G0}:$D{r},"Advance"))',
            NOTE, CALC_FILL, border=True)
        put(ws, f"L{r}",
            f'=IF(OR($I{r}="",$I{r}="?",$B{r}=""),"",'
            f'SUMPRODUCT(({WKS}!$A$6:$A${5+N_WEEKS}<=$J{r})*({WKS}!${{E0}}$4:${{E1}}$4=$I{r})*{WKS}!${{E0}}$6:${{E1}}${5+N_WEEKS})'
            f'-SUMIFS($E${G0}:$E${G1},$I${G0}:$I${G1},$I{r},$B${G0}:$B${G1},"<="&$B{r}))',
            CALC, CALC_FILL, MONEY_SIGNED, border=True)
        put(ws, f"M{r}",
            f'=IF($B{r}&$C{r}&$D{r}&$E{r}="","",'
            f'IF($B{r}="","missing date",IF($C{r}="","missing worker",'
            f'IF($I{r}="?","unknown worker - was the name changed on Workers?",'
            f'IF($D{r}="","missing type",IF(N($E{r})=0,"missing amount",'
            f'IF(AND($E{r}<0,$D{r}<>"Correction"),"negative amount - use a Correction to undo",'
            f'IF(OR($B{r}<{WKS}!$A$6,$B{r}>{WKS}!$B${5+N_WEEKS}),"date outside this workbook\'s year",'
            f'IF(AND($D{r}="Advance",N($L{r})<0),"Warning: more than earned so far - worker now owes "&ROUND(-$L{r},0),'
            f'""))))))))',
            CALC, CALC_FILL, border=True)
    flag_rule(ws, f"M{G0}:M{G1}", f"M{G0}")
    dv_date(ws, f"B{G0}:B{G1}")
    dv_list(ws, "=WorkerNames", f"C{G0}:C{G1}")
    dv_list(ws, "=EntryTypes", f"D{G0}:D{G1}",
            "Advance: money before payday. Weekly pay: payday cash. Correction: undo a mistake (may be negative).")
    dv_list(ws, "=YesNo", f"H{G0}:H{G1}")
    ws.freeze_panes = f"A{G0}"
    ws.auto_filter.ref = f"A5:M{G1}"

    # ============================================================ Weeks (engine)
    ws = ws_wks
    title(ws, "Weeks", "Calculation engine - nothing to type here. Every other sheet reads from it.")
    K0, K1 = 6, 6 + N_WEEKS - 1
    DAYS0 = 14                                # N: first days-worked column
    DAYS1 = DAYS0 + N_SLOTS - 1               # Y
    EARN0 = DAYS1 + 1                         # Z: first earnings column
    EARN1 = EARN0 + N_SLOTS - 1               # AK
    cSUM, cDIFF = L(EARN1 + 1), L(EARN1 + 2)
    lbl = ["Week of", "to", "Moulded", "Broken", "Good", "Paid blocks", "Pot (so'm)", "Crew-days",
           "Reject %", "Good per crew-day", "Cement bags", "Bags per 1000", "Check"]
    header_row(ws, 5, 1, lbl + [""] * (2 * N_SLOTS) + ["Shares total", "Shares - pot"],
               [11, 11, 10, 9, 10, 10, 13, 9, 8, 10, 9, 9, 30] + [8] * N_SLOTS + [11] * N_SLOTS + [13, 11],
               height=42)
    put(ws, f"{L(DAYS0)}3", "Days worked", H2)
    put(ws, f"{L(EARN0)}3", "Earned (so'm)", H2)
    for k in range(N_SLOTS):
        wid = f"W{k+1:02d}"
        for col in (DAYS0 + k, EARN0 + k):
            put(ws, f"{L(col)}4", wid, NOTE, align=CENTER)
            ws.cell(row=5, column=col).value = f'=IF({WRK}!$B${W0+k}="",{WRK}!$A${W0+k},{WRK}!$B${W0+k})'
    LG_WEEK = f"{LOG}!${cWEEK}${D0}:${cWEEK}${D1}"
    def log_col(c):
        return f"{LOG}!${c}${D0}:${c}${D1}"
    for i in range(N_WEEKS):
        r = K0 + i
        put(ws, f"A{r}", FIRST_DAY + dt.timedelta(days=7 * i), BOLD, fmt=DATE, align=CENTER, border=True)
        put(ws, f"B{r}", f"=$A{r}+6", CALC, fmt=DATE, align=CENTER, border=True)
        for col, src in (("C", "C"), ("D", "D"), ("E", "E"), ("F", "F"), ("G", cVAL), ("H", cCREW), ("K", "G")):
            put(ws, f"{col}{r}", f"=SUMIFS({log_col(src)},{LG_WEEK},$A{r})", CALC, CALC_FILL,
                MONEY if col == "G" else "# ##0.#", border=True)
        put(ws, f"I{r}", f'=IF($C{r}=0,"",$D{r}/$C{r})', CALC, CALC_FILL, PCT, CENTER, True)
        put(ws, f"J{r}", f'=IF($H{r}=0,"",$E{r}/$H{r})', CALC, CALC_FILL, "# ##0", CENTER, True)
        put(ws, f"L{r}", f'=IF(OR($E{r}=0,$K{r}=0),"",$K{r}/$E{r}*1000)', CALC, CALC_FILL, "0.0", CENTER, True)
        put(ws, f"M{r}",
            f'=IF(AND($G{r}>0,$H{r}=0),"pot not shared - no attendance marked",'
            f'IF(AND($I{r}<>"",N($I{r})>RejectAlarm),"reject rate high",'
            f'IF(AND($L{r}<>"",N(CementNorm)>0),IF(ABS($L{r}/CementNorm-1)>0.3,"cement use off by >30% - check counts",""),"")))',
            CALC, CALC_FILL, border=True)
        for k in range(N_SLOTS):
            dcol, ecol = L(DAYS0 + k), L(EARN0 + k)
            put(ws, f"{dcol}{r}", f"=SUMIFS({log_col(L(ATT0 + k))},{LG_WEEK},$A{r})",
                CALC, CALC_FILL, "0.#", CENTER, True)
            # whole so'm: cash is paid in whole so'm, and rounding here keeps
            # tiny fractions from turning into phantom debts on the statements
            put(ws, f"{ecol}{r}", f"=IF($H{r}=0,0,ROUND($G{r}*{dcol}{r}/$H{r},0))",
                CALC, CALC_FILL, MONEY, None, True)
        put(ws, f"{cSUM}{r}", f"=SUM({L(EARN0)}{r}:{L(EARN1)}{r})", CALC, CALC_FILL, MONEY, None, True)
        put(ws, f"{cDIFF}{r}", f"=${cSUM}{r}-$G{r}", CALC, CALC_FILL, MONEY_SIGNED, None, True)
    flag_rule(ws, f"M{K0}:M{K1}", f"M{K0}")
    name(wb, "WeekStarts", f"{WKS}!$A${K0}:$A${K1}")
    ws.freeze_panes = f"C{K0}"
    E0, E1 = L(EARN0), L(EARN1)

    # the ledger's balance formula needed the earnings columns; fill them in
    for i in range(N_LEDGER):
        c = ws_led[f"L{G0+i}"]
        c.value = c.value.replace("{E0}", E0).replace("{E1}", E1)

    WK_A = f"{WKS}!$A${K0}:$A${K1}"
    EARN_HDR = f"{WKS}!${E0}$4:${E1}$4"
    EARN_BLK = f"{WKS}!${E0}${K0}:${E1}${K1}"
    DAYS_HDR = f"{WKS}!${L(DAYS0)}$4:${L(DAYS1)}$4"
    DAYS_BLK = f"{WKS}!${L(DAYS0)}${K0}:${L(DAYS1)}${K1}"
    LE = lambda c: f"{LED}!${c}${G0}:${c}${G1}"

    def earned_before(idc, wk, op="<"):
        return f'SUMPRODUCT(({WK_A}{op}{wk})*({EARN_HDR}={idc})*{EARN_BLK})'
    def cash(idc, typ=None, lo=None, hi=None, before=None):
        parts = [f'{LE("E")}', f'{LE("I")},{idc}']
        if typ: parts.append(f'{LE("D")},"{typ}"')
        if lo:  parts.append(f'{LE("B")},">="&{lo}')
        if hi:  parts.append(f'{LE("B")},"<="&{hi}')
        if before: parts.append(f'{LE("B")},"<"&{before}')
        return "SUMIFS(" + ",".join(parts) + ")"
    def week_cell(block_blk, hdr, idc, wk):
        return f'IFERROR(INDEX({block_blk},MATCH({wk},{WK_A},0),MATCH({idc},{hdr},0)),0)'
    def wkinfo(col, wk):
        return f'IFERROR(INDEX({WKS}!${col}${K0}:${col}${K1},MATCH({wk},{WK_A},0)),0)'

    # ============================================================ Weekly Pay
    ws = ws_pay
    title(ws, "Weekly Pay", "Payday. Pick the week, pay each worker the green column, then record "
                            "each payment in the Cash Ledger as 'Weekly pay'.")
    put(ws, "B4", "Week starting (Monday):", BOLD, align=Alignment(horizontal="right"))
    wkc = "$C$4"
    put(ws, "C4", FIRST_DAY, Font(name=F, size=12, bold=True, color="0000FF"), INP_FILL, DATE, CENTER, True)
    dv_list(ws, "=WeekStarts", "C4", "Pick the Monday the week starts on.")
    put(ws, "D4", f'="to "&TEXT({wkc}+6,"dd.mm.yyyy")', BODY)
    # TEXT is used only for this caption; the maths never depends on it
    ws["D4"].value = f"={wkc}+6"
    ws["D4"].number_format = '"to "' + DATE

    summary = [("Good blocks", "E", "# ##0"), ("Paid blocks", "F", "# ##0"),
               ("Pot (so'm)", "G", MONEY), ("Crew-days", "H", "0.#"), ("Reject rate", "I", PCT)]
    for i, (lab, col, fmt) in enumerate(summary):
        cc = 2 + i * 2
        put(ws, f"{L(cc)}6", lab, NOTE, align=CENTER)
        put(ws, f"{L(cc)}7", f"={wkinfo(col, wkc)}", BIG, KEY_FILL, fmt, CENTER, True)
        ws.merge_cells(start_row=6, start_column=cc, end_row=6, end_column=cc + 1)
        ws.merge_cells(start_row=7, start_column=cc, end_row=7, end_column=cc + 1)
    put(ws, "L6", "Workbook checks", NOTE, align=CENTER)
    put(ws, "L7", f"={q(S_CHECKS)}!$C$4", BOLD, align=CENTER, border=True)
    ws.merge_cells("L6:N6"); ws.merge_cells("L7:N7")

    P0, P1 = 11, 11 + N_SLOTS - 1
    header_row(ws, 10, 1, ["ID", "Worker", "Status", "Days", "Share", "Earned this week",
                           "Brought forward", "Advances", "Corrections", "Due",
                           "TO PAY", "Paid (in ledger)", "Still to pay", "Carried forward", "Note"],
               [6, 20, 8, 7, 8, 14, 15, 13, 12, 14, 15, 14, 13, 15, 26], height=36)
    ws["K10"].fill = PatternFill("solid", fgColor="375623")
    wk_end = f"({wkc}+6)"
    for k in range(N_SLOTS):
        r = P0 + k
        idc = f"$A{r}"
        g = f'$B{r}=""'
        put(ws, f"A{r}", f"={WRK}!$A${W0+k}", NOTE, align=CENTER, border=True)
        put(ws, f"B{r}", f'=IF({WRK}!$B${W0+k}="","",{WRK}!$B${W0+k})', LINK, border=True)
        put(ws, f"C{r}", f'=IF({g},"",{WRK}!$C${W0+k})', LINK, align=CENTER, border=True)
        put(ws, f"D{r}", f'=IF({g},"",{week_cell(DAYS_BLK, DAYS_HDR, idc, wkc)})', CALC, CALC_FILL, "0.#", CENTER, True)
        put(ws, f"E{r}", f'=IF(OR({g},N($H$7)=0),"",$D{r}/$H$7)', CALC, CALC_FILL, PCT, CENTER, True)
        put(ws, f"F{r}", f'=IF({g},"",{week_cell(EARN_BLK, EARN_HDR, idc, wkc)})', CALC, CALC_FILL, MONEY, None, True)
        # brought forward: everything earned before this week minus all cash before it
        put(ws, f"G{r}", f'=IF({g},"",{earned_before(idc, wkc)}-{cash(idc, before=wkc)})',
            CALC, CALC_FILL, MONEY_SIGNED, None, True)
        put(ws, f"H{r}", f'=IF({g},"",{cash(idc, "Advance", wkc, wk_end)})', CALC, CALC_FILL, MONEY, None, True)
        put(ws, f"I{r}", f'=IF({g},"",{cash(idc, "Correction", wkc, wk_end)})', CALC, CALC_FILL, MONEY_SIGNED, None, True)
        put(ws, f"J{r}", f'=IF({g},"",$G{r}+$F{r}-$H{r}-$I{r})', CALC, CALC_FILL, MONEY_SIGNED, None, True)
        # never below zero; old debt comes back at most DebtCap of this week's earnings
        put(ws, f"K{r}",
            f'=IF({g},"",ROUND(MAX(0,IF($G{r}>=0,$J{r},$F{r}-$H{r}-$I{r}-MIN(-$G{r},DebtCap*$F{r}))),0))',
            Font(name=F, size=11, bold=True, color="375623"), OK_FILL, MONEY, None, True)
        put(ws, f"L{r}", f'=IF({g},"",{cash(idc, "Weekly pay", wkc, wk_end)})', CALC, CALC_FILL, MONEY, None, True)
        put(ws, f"M{r}", f'=IF({g},"",MAX(0,$K{r}-$L{r}))', CALC, CALC_FILL, MONEY, None, True)
        put(ws, f"N{r}", f'=IF({g},"",$J{r}-$L{r})', CALC, CALC_FILL, MONEY_SIGNED, None, True)
        put(ws, f"O{r}",
            f'=IF({g},"",IF(ROUND($N{r},0)<0,"owes you "&ROUND(-$N{r},0),'
            f'IF(AND(ROUND($N{r},0)>0,$M{r}>0),"not yet paid",'
            f'IF(ROUND($N{r},0)>0,"you owe "&ROUND($N{r},0),"settled"))))',
            CALC, CALC_FILL, border=True)
    tr = P1 + 1
    put(ws, f"B{tr}", "Crew total", BOLD, border=True)
    for col, fmt in (("D", "0.#"), ("F", MONEY), ("G", MONEY_SIGNED), ("H", MONEY), ("I", MONEY_SIGNED),
                     ("J", MONEY_SIGNED), ("K", MONEY), ("L", MONEY), ("M", MONEY), ("N", MONEY_SIGNED)):
        put(ws, f"{col}{tr}", f"=SUM({col}{P0}:{col}{P1})", BOLD, KEY_FILL, fmt, None, True)

    ck = tr + 2
    put(ws, f"B{ck}", "Before you pay", H2)
    checks = [
        ("Shares add up to the pot", f'=IF(ABS($F${tr}-$F$7)<={N_SLOTS},"OK","MISMATCH: "&ROUND($F${tr}-$F$7,0))'),
        ("Everyone's days match the crew total", f'=IF(ABS($D${tr}-$H$7)<0.001,"OK","MISMATCH - check attendance")'),
        ("Pot is shared out", f'=IF(AND($F$7>0,$H$7=0),"NO - mark attendance in the Daily Log","OK")'),
        ("Reject rate", f'=IF($F$7=0,"-",IF(N($J$7)>RejectAlarm,"HIGH - look at curing and handling","OK"))'),
        ("Cash needed for payday", f"=$M${tr}"),
    ]
    for i, (lab, f) in enumerate(checks):
        r = ck + 1 + i
        put(ws, f"B{r}", lab, BODY)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = put(ws, f"F{r}", f, BOLD, border=True)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        if i == 4:
            c.number_format = MONEY
    ws.conditional_formatting.add(f"F{ck+1}:F{ck+4}",
        FormulaRule(formula=[f'AND(F{ck+1}<>"OK",F{ck+1}<>"-")'], fill=WARN_FILL))
    ws.conditional_formatting.add("L7", FormulaRule(formula=['L7<>"ALL CHECKS OK"'], fill=WARN_FILL))
    ws.conditional_formatting.add("L7", FormulaRule(formula=['L7="ALL CHECKS OK"'], fill=OK_FILL))

    notes = [
        "Brought forward: positive = you still owe the worker from earlier weeks; negative = the worker owes you.",
        "TO PAY is never below zero. If advances were bigger than earnings, pay is 0 and the rest carries to next week.",
        "After paying, record each payment in the Cash Ledger as 'Weekly pay' with a date inside this week - "
        "'Still to pay' then drops to 0.",
    ]
    for i, n in enumerate(notes):
        put(ws, f"B{ck+7+i}", n, NOTE)
    ws.freeze_panes = "C11"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ============================================================ Payslip
    ws = ws_slip
    title(ws, "Payslip", "Pick a week and a worker, print, and have the worker sign when paid.")
    for col, w in zip("ABCDEF", [2, 30, 18, 18, 26, 4]):
        ws.column_dimensions[col].width = w
    put(ws, "B4", "Week starting (Monday)", BOLD)
    put(ws, "C4", FIRST_DAY, INP, INP_FILL, DATE, CENTER, True)
    dv_list(ws, "=WeekStarts", "C4")
    put(ws, "B5", "Worker", BOLD)
    put(ws, "C5", f"={WRK}!$B${W0}", INP, INP_FILL, border=True)
    dv_list(ws, "=WorkerNames", "C5")
    put(ws, "D5", '=IFERROR(INDEX(WorkerIDs,MATCH($C$5,WorkerNames,0)),"unknown worker")', NOTE)
    idc, wk = "$D$5", "$C$4"

    put(ws, "B7", "DONABAY - PAYSLIP", Font(name=F, size=14, bold=True, color="1F3864"))
    put(ws, "B8", '=$C$5', BOLD)
    put(ws, "C8", f"={wk}", BODY, fmt=DATE)
    put(ws, "D8", f"={wk}+6", BODY, fmt='"to "' + DATE)

    lines = [
        (10, "Crew paid blocks this week", f"={wkinfo('F', wk)}", "# ##0"),
        (11, "Crew pot", f"={wkinfo('G', wk)}", MONEY),
        (12, "Crew-days worked", f"={wkinfo('H', wk)}", "0.#"),
        (13, "Your days", f"={week_cell(DAYS_BLK, DAYS_HDR, idc, wk)}", "0.#"),
        (14, "Your share", f'=IF(N($C$12)=0,0,$C$13/$C$12)', PCT),
        (16, "Earned this week", f"={week_cell(EARN_BLK, EARN_HDR, idc, wk)}", MONEY),
        (17, "Brought forward from earlier weeks", f"={earned_before(idc, wk)}-{cash(idc, before=wk)}", MONEY_SIGNED),
        (18, "Advances this week (listed below)", f"=-{cash(idc, 'Advance', wk, '(' + wk + '+6)')}", MONEY_SIGNED),
        (19, "Corrections this week", f"=-{cash(idc, 'Correction', wk, '(' + wk + '+6)')}", MONEY_SIGNED),
        (20, "Due", "=$C$16+$C$17+$C$18+$C$19", MONEY_SIGNED),
        (22, "TO PAY NOW", "=ROUND(MAX(0,IF($C$17>=0,$C$20,$C$16+$C$18+$C$19-MIN(-$C$17,DebtCap*$C$16))),0)", MONEY),
        (23, "Already paid this week (ledger)", f"={cash(idc, 'Weekly pay', wk, '(' + wk + '+6)')}", MONEY),
        (24, "Carried to next week", "=$C$20-$C$23", MONEY_SIGNED),
    ]
    for r, lab, f, fmt in lines:
        big = r in (22,)
        put(ws, f"B{r}", lab, BOLD if r in (16, 20, 22, 24) else BODY)
        put(ws, f"C{r}", f, Font(name=F, size=12, bold=True, color="375623") if big else BOLD if r in (20, 24) else CALC,
            OK_FILL if big else CALC_FILL, fmt, None, True)
    put(ws, "D24", '=IF(ROUND($C$24,0)<0,"you owe the workshop "&ROUND(-$C$24,0)&" so\'m",'
                   'IF(ROUND($C$24,0)>0,"the workshop owes you "&ROUND($C$24,0)&" so\'m","settled"))', NOTE)

    put(ws, "B26", "Advances taken this week", H2)
    header_row(ws, 27, 2, ["Date", "Amount (so'm)", "Reason"], height=20)
    for n in range(1, N_SLIP_ADV + 1):
        r = 27 + n
        key = f'{idc}&"|"&{wk}&"|"&{n}'
        put(ws, f"B{r}", f'=IFERROR(INDEX({LE("B")},MATCH({key},{LE("K")},0)),"")', CALC, fmt=DATE, border=True)
        put(ws, f"C{r}", f'=IFERROR(INDEX({LE("E")},MATCH({key},{LE("K")},0)),"")', CALC, fmt=MONEY, border=True)
        put(ws, f"D{r}", f'=IFERROR(INDEX({LE("F")},MATCH({key},{LE("K")},0)),"")', CALC, border=True)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    more = 28 + N_SLIP_ADV
    put(ws, f"B{more}",
        f'=IF(COUNTIFS({LE("I")},{idc},{LE("D")},"Advance",{LE("B")},">="&{wk},{LE("B")},"<="&({wk}+6))>{N_SLIP_ADV},'
        f'"More advances than fit here - see the Cash Ledger","")', NOTE)
    sig = more + 2
    for i, lab in enumerate(["Worker's signature", "Paid by", "Date paid"]):
        put(ws, f"B{sig + 2*i}", lab, BODY)
        c = ws[f"C{sig + 2*i}"]
        c.border = Border(bottom=Side(style="thin", color="000000"))
        ws.merge_cells(start_row=sig + 2*i, start_column=3, end_row=sig + 2*i, end_column=5)
    ws.print_area = f"B7:E{sig + 5}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ============================================================ Worker History
    ws = ws_hist
    title(ws, "Worker History", "One worker's whole account, week by week: earned, taken, paid, and "
                                "where the balance stands.")
    put(ws, "B4", "Worker", BOLD)
    put(ws, "C4", f"={WRK}!$B${W0}", INP, INP_FILL, border=True)
    dv_list(ws, "=WorkerNames", "C4")
    put(ws, "D4", '=IFERROR(INDEX(WorkerIDs,MATCH($C$4,WorkerNames,0)),"unknown worker")', NOTE)
    hid = "$D$4"
    H0, H1 = 7, 7 + N_WEEKS - 1
    header_row(ws, 6, 1, ["Week of", "Days", "Earned", "Advances", "Corrections", "Weekly pay",
                          "Balance after the week", "Meaning"],
               [12, 7, 14, 13, 12, 13, 18, 30])
    put(ws, "G5", "+ you owe them / - they owe you", NOTE, align=CENTER)
    for i in range(N_WEEKS):
        r = H0 + i
        wkr = f"$A{r}"
        put(ws, f"A{r}", f"={WKS}!$A${K0+i}", BODY, fmt=DATE, align=CENTER, border=True)
        live = f"{wkr}>TODAY()"                          # future weeks stay blank
        put(ws, f"B{r}", f'=IF({live},"",{week_cell(DAYS_BLK, DAYS_HDR, hid, wkr)})', CALC, CALC_FILL, "0.#", CENTER, True)
        put(ws, f"C{r}", f'=IF({live},"",{week_cell(EARN_BLK, EARN_HDR, hid, wkr)})', CALC, CALC_FILL, MONEY, None, True)
        for col, typ in (("D", "Advance"), ("E", "Correction"), ("F", "Weekly pay")):
            put(ws, f"{col}{r}", f'=IF({live},"",{cash(hid, typ, wkr, "(" + wkr + "+6)")})',
                CALC, CALC_FILL, MONEY_SIGNED if typ == "Correction" else MONEY, None, True)
        # computed directly, not by chaining the row above, so one bad row
        # cannot poison every week after it
        put(ws, f"G{r}", f'=IF({live},"",{earned_before(hid, wkr, "<=")}-{cash(hid, hi="(" + wkr + "+6)")})',
            BOLD, CALC_FILL, MONEY_SIGNED, None, True)
        put(ws, f"H{r}", f'=IF(OR({live},$G{r}=""),"",IF(ROUND($G{r},0)<0,"owes you",'
                         f'IF(ROUND($G{r},0)>0,"you owe","settled")))', CALC, CALC_FILL, border=True)
    ws.freeze_panes = f"A{H0}"

    # ============================================================ Checks
    ws = ws_chk
    title(ws, "Checks", "The workbook checking itself. Everything should say OK before payday.")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 50
    put(ws, "B4", "Overall", H2)
    header_row(ws, 6, 2, ["Check", "Result", "Where to look"], height=22)
    LOGCHK = f"{LOG}!${cCHK}${D0}:${cCHK}${D1}"
    items = [
        ("Worker list is clean (names unique, status set)",
         f'=COUNTIF({WRK}!$H${W0}:$H${W1},"?*")', S_WORKERS),
        ("Daily Log rows without problems",
         f'=COUNTIF({LOGCHK},"?*")', f"{S_LOG} - Check column"),
        ("Cash Ledger entries are complete and valid",
         f'=COUNTIF({LE("M")},"?*")-COUNTIF({LE("M")},"Warning*")', f"{S_LEDGER} - Check column"),
        ("Every week's pot is shared out",
         f'=COUNTIF({WKS}!$M${K0}:$M${K1},"pot*")', S_WEEKS),
        ("Worker shares add up to each week's pot",
         f'=SUMPRODUCT(--(ABS({WKS}!${cDIFF}${K0}:${cDIFF}${K1})>{N_SLOTS}))', S_WEEKS),
        ("A pay rate is set",
         f'=IF(COUNTIF(RateValues,">0")=0,1,0)', S_SET),
        ("Rate dates are in order, oldest first",
         f'=SUMPRODUCT(({SET}!$B${R0+1}:$B${R1}<>"")*({SET}!$B${R0+1}:$B${R1}<={SET}!$B${R0}:$B${R1-1}))', S_SET),
        ("Cash Ledger has room left",
         f'=IF(COUNT({LE("B")})>{N_LEDGER - 50},1,0)', f"{S_LEDGER} - nearly full, ask for a rebuild"),
        ("Worker slots left",
         f'=IF(COUNTA({WRK}!$B${W0}:$B${W1})>={N_SLOTS},1,0)', f"{S_WORKERS} - all slots used, ask for a rebuild"),
    ]
    C0 = 7
    for i, (lab, f, where) in enumerate(items):
        r = C0 + i
        put(ws, f"B{r}", lab, BODY, border=True)
        put(ws, f"C{r}", f'=IF(({f[1:]})=0,"OK",({f[1:]})&" to fix")', BOLD, align=CENTER, border=True)
        put(ws, f"D{r}", where, NOTE, border=True)
    CN = C0 + len(items) - 1
    put(ws, f"B{CN+2}", "Warnings (not errors)", H2)
    put(ws, f"B{CN+3}", "Advances bigger than what the worker had earned so far", BODY, border=True)
    put(ws, f"C{CN+3}", f'=COUNTIF({LE("M")},"Warning*")', BOLD, align=CENTER, border=True)
    put(ws, f"D{CN+3}", "Allowed - just so you know. Shows as debt carried forward.", NOTE, border=True)
    put(ws, f"B{CN+4}", "Weeks with a high reject rate or odd cement use", BODY, border=True)
    put(ws, f"C{CN+4}", f'=COUNTIF({WKS}!$M${K0}:$M${K1},"?*")-COUNTIF({WKS}!$M${K0}:$M${K1},"pot*")',
        BOLD, align=CENTER, border=True)
    put(ws, f"D{CN+4}", S_WEEKS, NOTE, border=True)
    put(ws, "C4", f'=IF(COUNTIF($C${C0}:$C${CN},"OK")={len(items)},"ALL CHECKS OK",'
                  f'({len(items)}-COUNTIF($C${C0}:$C${CN},"OK"))&" check(s) need attention")',
        Font(name=F, size=12, bold=True), border=True)
    ws.merge_cells("C4:D4")
    ws.conditional_formatting.add(f"C{C0}:C{CN}", FormulaRule(formula=[f'C{C0}<>"OK"'], fill=WARN_FILL))
    ws.conditional_formatting.add(f"C{C0}:C{CN}", FormulaRule(formula=[f'C{C0}="OK"'], fill=OK_FILL))
    ws.conditional_formatting.add("C4", FormulaRule(formula=['C4="ALL CHECKS OK"'], fill=OK_FILL))
    ws.conditional_formatting.add("C4", FormulaRule(formula=['C4<>"ALL CHECKS OK"'], fill=WARN_FILL))

    # ============================================================ Start here
    ws = ws_start
    title(ws, "Donabay", "Filler-block production and weekly crew pay. No macros, no add-ins - just type.")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 92
    put(ws, "B4", "Workbook status", BOLD)
    put(ws, "C4", f"={q(S_CHECKS)}!$C$4", BOLD, border=True)
    ws.conditional_formatting.add("C4", FormulaRule(formula=['C4="ALL CHECKS OK"'], fill=OK_FILL))
    ws.conditional_formatting.add("C4", FormulaRule(formula=['C4<>"ALL CHECKS OK"'], fill=WARN_FILL))
    steps = [
        ("FIRST TIME", None),
        ("1. Settings", "Type your pay rate per block next to the start date."),
        ("2. Workers", "Replace 'Worker 1..4' with real names. Leave them Active."),
        ("EVERY DAY  (1 minute)", None),
        ("Daily Log", "On today's row: blocks moulded, blocks broken, and 1 under each worker who came "
                      "(0.5 for half a day, leave empty if absent). Broken blocks can be filled in later, "
                      "once you know them."),
        ("WHEN SOMEONE ASKS FOR MONEY", None),
        ("Cash Ledger", "New row: date, worker, type 'Advance', amount, reason. That's it - it is subtracted "
                        "on payday automatically."),
        ("PAYDAY", None),
        ("1. Weekly Pay", "Pick the week. Check 'Before you pay' says OK. Pay each worker the green TO PAY amount."),
        ("2. Cash Ledger", "Record each payment as type 'Weekly pay', dated inside that week."),
        ("3. Payslip", "Print one per worker and have them sign. It lists every advance they took that week."),
        ("PEOPLE CHANGE", None),
        ("Someone joins", "Workers: type the name in the next empty slot, status Active, joined date."),
        ("Someone leaves", "Workers: set status Left. Do not delete the row - their history stays."),
        ("Rate changes", "Settings: add a NEW row with the date the new rate starts. Old weeks keep the old rate."),
        ("MISTAKES", None),
        ("Wrong ledger entry", "Do not delete it. Add a 'Correction' row with the opposite amount "
                               "(e.g. -100000) and say why."),
        ("RULES THIS FILE FOLLOWS", None),
        ("Pay base", "Only good blocks are paid (moulded minus broken). Settings can add a breakage allowance."),
        ("Split", "The week's pot is shared by days worked: pot x your days / everyone's days."),
        ("Debts", "Advances bigger than earnings never make pay negative - pay is 0 and the rest carries "
                  "forward, visible on every payslip until it is paid back."),
        ("Colours", "Yellow = you type here.  Grey = calculated, leave alone.  Pink = needs a look."),
    ]
    r = 6
    for a, b in steps:
        if b is None:
            r += 1
            put(ws, f"B{r}", a, H2)
        else:
            put(ws, f"B{r}", a, BOLD, align=Alignment(vertical="top"))
            put(ws, f"C{r}", b, BODY, align=WRAP)
            ws.row_dimensions[r].height = 28 if len(b) > 95 else 15
        r += 1

    # ------------------------------------------------------------ demo data
    if demo:
        fill_demo(wb, today_row, D0, ATT0, G0, R0, W0, week, cap)
    return wb


def fill_demo(wb, today_row, D0, ATT0, G0, R0, W0, week=None, cap=None):
    """Synthetic data for verification only. Never shipped."""
    s, w, lg = wb[S_SET], wb[S_WORKERS], wb[S_LEDGER]
    wb[S_SLIP]["C5"] = "Bobur"
    wb[S_HIST]["C4"] = "Davron"
    if week:
        wb[S_PAY]["C4"] = week
        wb[S_SLIP]["C4"] = week
    if cap is not None:
        s["C6"] = cap
    s[f"C{R0}"] = 200                                   # 200 so'm from 21 Sep
    s[f"B{R0+1}"] = dt.date(2026, 9, 26)                # rate rises on Saturday
    s[f"C{R0+1}"] = 250
    for i, nm in enumerate(["Ali", "Bobur", "Davron", "Eldor"]):
        w[f"B{W0+i}"] = nm
    w[f"B{W0+4}"] = "Farhod"; w[f"C{W0+4}"] = "Left"
    log = wb[S_LOG]
    days = {  # date: (moulded, broken, [Ali, Bobur, Davron, Eldor])
        dt.date(2026, 9, 21): (820, 20, [1, 1, 1, 1]),
        dt.date(2026, 9, 22): (800, 10, [1, 1, None, 1]),
        dt.date(2026, 9, 23): (840, 40, [1, 0.5, 1, 1]),
        dt.date(2026, 9, 24): (810, 5, [1, 1, 1, 1]),
        dt.date(2026, 9, 25): (790, 15, [1, 1, 1, 1]),
        dt.date(2026, 9, 26): (700, 30, [1, 1, 1, 1]),
        dt.date(2026, 9, 28): (800, 10, [1, 1, 1, 1]),
    }
    for d, (m, b, att) in days.items():
        r = D0 + (d - FIRST_DAY).days
        log[f"C{r}"], log[f"D{r}"] = m, b
        for k, v in enumerate(att):
            if v is not None:
                log.cell(row=r, column=ATT0 + k, value=v)
    entries = [
        (dt.date(2026, 9, 23), "Bobur", "Advance", 150000, "medicine"),
        (dt.date(2026, 9, 24), "Ali", "Advance", 50000, "transport"),
        (dt.date(2026, 9, 25), "Bobur", "Advance", 100000, "family"),
        (dt.date(2026, 9, 26), "Davron", "Advance", 2000000, "big loan"),
        (dt.date(2026, 9, 24), "Eldor", "Advance", 30000, "entered by mistake"),
        (dt.date(2026, 9, 24), "Eldor", "Correction", -30000, "undo the mistaken advance"),
        (dt.date(2026, 9, 26), "Ali", "Weekly pay", 400000, "part of the week's pay"),
    ]
    for i, (d, who, typ, amt, why) in enumerate(entries):
        r = G0 + i
        lg[f"B{r}"], lg[f"C{r}"], lg[f"D{r}"], lg[f"E{r}"], lg[f"F{r}"] = d, who, typ, amt, why


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--demo":
        out = sys.argv[2]
        week = dt.date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else None
        cap = float(sys.argv[4]) if len(sys.argv) > 4 else None
        build(demo=True, week=week, cap=cap).save(out)
    else:
        out = os.path.join(ROOT, "Donabay.xlsx")
        build().save(out)
    print("wrote", out)
